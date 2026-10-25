import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch

from googleapiclient.errors import HttpError
from openpyxl import Workbook, load_workbook

from job_tools import drive_export


def prospect(key, title):
    return {"job_key": key, "job_title": title, "url": f"https://example.com/{key}"}


class DriveExportTests(unittest.TestCase):
    @staticmethod
    def existing_workbook(with_tab=True):
        workbook = Workbook()
        workbook.active.title = "Overview"
        workbook.active["A1"] = "Keep this value"
        if with_tab:
            sheet = workbook.create_sheet("Prospective jobs")
            sheet.append(drive_export.HEADER)
            sheet.append(drive_export.prospect_rows([prospect("old", "Old job")])[1])
            sheet["M2"] = "Manual note"
        workbook.create_sheet("Notes")["B2"] = "Keep this too"
        data = BytesIO()
        workbook.save(data)
        return data.getvalue()

    def test_workbook_appends_only_new_jobs_and_preserves_existing_cells(self):
        rows = drive_export.prospect_rows([
            prospect("old", "Changed title"),
            prospect("new", '=HYPERLINK("https://example.com")'),
            prospect("new", "Duplicate in input"),
        ])
        data, added = drive_export.workbook_bytes(rows, self.existing_workbook())
        workbook = load_workbook(data)
        self.assertEqual(added, 1)
        self.assertEqual(workbook.sheetnames, ["Overview", "Prospective jobs", "Notes"])
        self.assertEqual(workbook["Overview"]["A1"].value, "Keep this value")
        self.assertEqual(workbook["Notes"]["B2"].value, "Keep this too")
        sheet = workbook["Prospective jobs"]
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet["A2"].value, "Old job")
        self.assertEqual(sheet["M2"].value, "Manual note")
        self.assertEqual(sheet["A3"].value, '=HYPERLINK("https://example.com")')
        self.assertEqual(sheet["A3"].data_type, "s")
        repeated, repeated_added = drive_export.workbook_bytes(rows, data.getvalue())
        self.assertEqual(repeated_added, 0)
        self.assertEqual(load_workbook(repeated)["Prospective jobs"].max_row, 3)

    def test_workbook_adds_tab_when_missing(self):
        data, added = drive_export.workbook_bytes(
            drive_export.prospect_rows([prospect("new", "Analyst")]),
            self.existing_workbook(with_tab=False),
        )
        workbook = load_workbook(data)
        self.assertEqual(added, 1)
        self.assertEqual(workbook.sheetnames, ["Overview", "Notes", "Prospective jobs"])
        self.assertEqual(workbook["Prospective jobs"]["A1"].value, "Job title")
        self.assertEqual(workbook["Prospective jobs"]["A2"].value, "Analyst")

    def test_unexpected_header_refuses_to_change_workbook(self):
        workbook = load_workbook(BytesIO(self.existing_workbook()))
        workbook["Prospective jobs"]["A1"] = "Different column"
        data = BytesIO()
        workbook.save(data)
        with self.assertRaisesRegex(ValueError, "unexpected columns"):
            drive_export.workbook_bytes(drive_export.prospect_rows([prospect("new", "New")]), data.getvalue())

    def test_upload_skips_excel_update_when_no_new_jobs(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "prospective_jobs.json"
            source.write_text(json.dumps([prospect("old", "Changed")]), encoding="utf-8")
            service = MagicMock()
            service.files.return_value.get.return_value.execute.return_value = {
                "id": "target", "name": "jobs.xlsx", "mimeType": drive_export.XLSX_MIME
            }
            service.files.return_value.get_media.return_value.execute.return_value = self.existing_workbook()
            with patch.object(drive_export, "drive_credentials", return_value=object()), \
                 patch.object(drive_export, "build", return_value=service):
                added, target = drive_export.upload_json(source, "target", Path("client.json"), Path("token.json"))
            self.assertEqual(added, 0)
            self.assertEqual(target["id"], "target")
            service.files.return_value.update.assert_not_called()

    def test_upload_creates_spreadsheet_when_no_file_id(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "prospective_jobs.json"
            source.write_text(json.dumps([prospect("new", "Analyst")]), encoding="utf-8")
            drive = MagicMock()
            drive.files.return_value.create.return_value.execute.return_value = {
                "id": "created", "name": "Job Watcher", "mimeType": drive_export.SHEETS_MIME
            }
            with patch.object(drive_export, "drive_credentials", return_value=object()), \
                 patch.object(drive_export, "build", return_value=drive), \
                 patch.object(drive_export, "update_google_sheet", return_value=(33, 1)) as update:
                created_ids = []
                added, target = drive_export.upload_json(
                    source, "", Path("client.json"), Path("token.json"), on_create=created_ids.append
                )
            self.assertEqual(added, 1)
            self.assertEqual(target["id"], "created")
            self.assertIn("/created/edit#gid=33", target["webViewLink"])
            self.assertEqual(drive.files.return_value.create.call_args.kwargs["body"], {
                "name": "Job Watcher", "mimeType": drive_export.SHEETS_MIME
            })
            update.assert_called_once()
            self.assertEqual(created_ids, ["created"])

    def test_upload_recreates_missing_configured_spreadsheet(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "prospective_jobs.json"
            source.write_text("[]", encoding="utf-8")
            drive = MagicMock()
            response = MagicMock(status=404, reason="Not Found")
            drive.files.return_value.get.return_value.execute.side_effect = HttpError(response, b"Not Found")
            drive.files.return_value.create.return_value.execute.return_value = {
                "id": "replacement", "name": "Job Watcher", "mimeType": drive_export.SHEETS_MIME
            }
            with patch.object(drive_export, "drive_credentials", return_value=object()), \
                 patch.object(drive_export, "build", return_value=drive), \
                 patch.object(drive_export, "update_google_sheet", return_value=(33, 0)):
                created_ids = []
                added, target = drive_export.upload_json(
                    source, "missing", Path("client.json"), Path("token.json"), on_create=created_ids.append
                )
            self.assertEqual(added, 0)
            self.assertEqual(target["id"], "replacement")
            self.assertEqual(created_ids, ["replacement"])

    def test_native_google_sheet_appends_only_new_jobs(self):
        sheets = MagicMock()
        spreadsheets = sheets.spreadsheets.return_value
        spreadsheets.get.return_value.execute.return_value = {
            "sheets": [{"properties": {"sheetId": 22, "title": "Prospective jobs"}}]
        }
        spreadsheets.values.return_value.get.return_value.execute.return_value = {
            "values": drive_export.prospect_rows([prospect("old", "Old job")])
        }
        rows = drive_export.prospect_rows([prospect("old", "Changed"), prospect("new", "=not a formula")])
        with patch.object(drive_export, "build", return_value=sheets):
            sheet_id, added = drive_export.update_google_sheet(object(), "target", rows)
        self.assertEqual((sheet_id, added), (22, 1))
        spreadsheets.batchUpdate.assert_not_called()
        spreadsheets.values.return_value.update.assert_not_called()
        kwargs = spreadsheets.values.return_value.append.call_args.kwargs
        self.assertEqual(kwargs["valueInputOption"], "RAW")
        self.assertEqual(kwargs["insertDataOption"], "INSERT_ROWS")
        self.assertEqual(kwargs["body"]["values"][0][0], "=not a formula")
        self.assertEqual(kwargs["body"]["values"][0][-1], "new")

    def test_native_google_sheet_creates_tab_if_missing(self):
        sheets = MagicMock()
        spreadsheets = sheets.spreadsheets.return_value
        spreadsheets.get.return_value.execute.return_value = {
            "sheets": [{"properties": {"sheetId": 11, "title": "Overview"}}]
        }
        spreadsheets.batchUpdate.return_value.execute.return_value = {
            "replies": [{"addSheet": {"properties": {"sheetId": 33}}}]
        }
        with patch.object(drive_export, "build", return_value=sheets):
            sheet_id, added = drive_export.update_google_sheet(
                object(), "target", drive_export.prospect_rows([prospect("new", "Analyst")])
            )
        self.assertEqual((sheet_id, added), (33, 1))
        self.assertEqual(spreadsheets.batchUpdate.call_args.kwargs["body"]["requests"][0]
                         ["addSheet"]["properties"]["title"], "Prospective jobs")
        self.assertEqual(spreadsheets.values.return_value.update.call_args.kwargs["body"],
                         {"values": [drive_export.HEADER]})
        self.assertEqual(spreadsheets.values.return_value.append.call_args.kwargs["body"]
                         ["values"][0][-1], "new")

    def test_save_new_file_id_keeps_other_settings(self):
        with tempfile.TemporaryDirectory() as directory:
            env_file = Path(directory) / ".env"
            env_file.write_text("GOOGLE_DRIVE_FILE_ID=\nJOB_SEARCH_LOCATION=nyc\n", encoding="utf-8")
            drive_export._save_file_id(env_file, "new-id")
            self.assertEqual(env_file.read_text(encoding="utf-8"),
                             "GOOGLE_DRIVE_FILE_ID=new-id\nJOB_SEARCH_LOCATION=nyc\n")


if __name__ == "__main__":
    unittest.main()
