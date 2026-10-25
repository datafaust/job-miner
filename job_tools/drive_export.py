"""Append new prospective jobs to a Google Sheet or Excel workbook on Drive."""

import argparse
import json
import os
from io import BytesIO
from pathlib import Path

from dotenv import load_dotenv
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaIoBaseUpload
from openpyxl import load_workbook
from openpyxl.styles import Font

from job_tools.job_stats import source_counts


PROJECT_DIR = Path(__file__).resolve().parent.parent
SCOPES = ["https://www.googleapis.com/auth/drive"]
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEETS_MIME = "application/vnd.google-apps.spreadsheet"
SHEET_NAME = "Prospective jobs"
SPREADSHEET_NAME = "Job Watcher"
COLUMNS = (
    ("job_title", "Job title"),
    ("company", "Company"),
    ("salary_range", "Salary range"),
    ("posted_at", "Posted at"),
    ("posted_text", "Posted text"),
    ("remote_status", "Remote status"),
    ("location", "Location"),
    ("url", "URL"),
    ("source", "Source"),
    ("source_saved_at", "Source saved at"),
    ("extracted_at", "Extracted at"),
    ("job_key", "Job key"),
)
HEADER = [label for _, label in COLUMNS]
URL_COLUMN = next(i for i, (key, _) in enumerate(COLUMNS) if key == "url")
KEY_COLUMN = next(i for i, (key, _) in enumerate(COLUMNS) if key == "job_key")


def prospect_rows(prospects):
    if not isinstance(prospects, list) or any(not isinstance(row, dict) for row in prospects):
        raise ValueError("Prospective jobs JSON must be an array of objects")
    return [HEADER] + [
        ["" if prospect.get(key) is None else str(prospect[key]) for key, _ in COLUMNS]
        for prospect in prospects
    ]


def _identity(row):
    """Use the stable job key, falling back to URL for older exports."""
    key = str(row[KEY_COLUMN]).strip() if len(row) > KEY_COLUMN and row[KEY_COLUMN] else ""
    url = str(row[URL_COLUMN]).strip() if len(row) > URL_COLUMN and row[URL_COLUMN] else ""
    return key, url


def _new_rows(incoming, existing):
    if existing and list(existing[0][:len(HEADER)]) != HEADER:
        raise ValueError(f"{SHEET_NAME} has unexpected columns; no rows were changed")
    keys = set()
    urls = set()
    for row in existing[1:]:
        key, url = _identity(row)
        if key:
            keys.add(key)
        if url:
            urls.add(url)
    additions = []
    for row in incoming[1:]:
        key, url = _identity(row)
        if not key and not url:
            raise ValueError("Each exported job needs a job_key or URL")
        if (key and key in keys) or (url and url in urls):
            continue
        additions.append(row)
        if key:
            keys.add(key)
        if url:
            urls.add(url)
    return additions


def _format_sheet(sheet):
    sheet.append(HEADER)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for column, width in {"A": 40, "B": 28, "C": 25, "D": 27, "E": 27,
                          "F": 16, "G": 35, "H": 65, "I": 18, "J": 27,
                          "K": 27, "L": 40}.items():
        sheet.column_dimensions[column].width = width


def workbook_bytes(rows, existing_content):
    """Return updated workbook data and the number of rows added."""
    workbook = load_workbook(BytesIO(existing_content))
    sheet = next(
        (item for item in workbook if item.title.casefold() == SHEET_NAME.casefold()),
        None,
    )
    if sheet is None:
        sheet = workbook.create_sheet(SHEET_NAME)
        _format_sheet(sheet)
        existing = [HEADER]
    else:
        existing = [list(row) for row in sheet.iter_rows(values_only=True)]
    additions = _new_rows(rows, existing)
    for row in additions:
        sheet.append([None] * len(COLUMNS))
        for column, value in enumerate(row, start=1):
            if value:
                cell = sheet.cell(sheet.max_row, column, value)
                cell.data_type = "s"  # Treat posting text as literal, even if it starts with '='.
    sheet.auto_filter.ref = sheet.dimensions
    data = BytesIO()
    workbook.save(data)
    data.seek(0)
    return data, len(additions)


def update_google_sheet(credentials, file_id, rows):
    """Create the named tab if needed and append only unseen jobs."""
    spreadsheets = build("sheets", "v4", credentials=credentials).spreadsheets()
    metadata = spreadsheets.get(
        spreadsheetId=file_id, fields="sheets(properties(sheetId,title))"
    ).execute()
    existing = next(
        (sheet["properties"] for sheet in metadata.get("sheets", [])
         if sheet["properties"]["title"].casefold() == SHEET_NAME.casefold()),
        None,
    )
    if existing is None:
        response = spreadsheets.batchUpdate(
            spreadsheetId=file_id,
            body={"requests": [{"addSheet": {"properties": {
                "title": SHEET_NAME,
                "gridProperties": {"frozenRowCount": 1},
            }}}]},
        ).execute()
        sheet_id = response["replies"][0]["addSheet"]["properties"]["sheetId"]
        existing_rows = []
        tab_title = SHEET_NAME
    else:
        sheet_id = existing["sheetId"]
        tab_title = existing["title"]
        existing_rows = spreadsheets.values().get(
            spreadsheetId=file_id,
            range=f"'{tab_title}'!A:L",
            valueRenderOption="FORMATTED_VALUE",
        ).execute().get("values", [])
    additions = _new_rows(rows, existing_rows or [HEADER])
    values = spreadsheets.values()
    if not existing_rows:
        values.update(
            spreadsheetId=file_id,
            range=f"'{SHEET_NAME}'!A1:L1",
            valueInputOption="RAW",
            body={"values": [HEADER]},
        ).execute()
    if additions:
        values.append(
            spreadsheetId=file_id,
            range=f"'{tab_title}'!A:L",
            valueInputOption="RAW",
            insertDataOption="INSERT_ROWS",
            body={"values": additions},
        ).execute()
    return sheet_id, len(additions)


def drive_credentials(client_secrets_file, token_file):
    credentials = None
    if token_file.exists():
        credentials = Credentials.from_authorized_user_file(str(token_file), SCOPES)
    if not credentials or not credentials.valid:
        if credentials and credentials.expired and credentials.refresh_token:
            credentials.refresh(Request())
        else:
            credentials = InstalledAppFlow.from_client_secrets_file(
                str(client_secrets_file), SCOPES
            ).run_local_server(port=0)
        token_file.parent.mkdir(parents=True, exist_ok=True)
        token_file.write_text(credentials.to_json(), encoding="utf-8")
        token_file.chmod(0o600)
    return credentials


def upload_json(input_file, file_id, client_secrets_file, token_file, on_create=None):
    prospects = json.loads(input_file.read_text(encoding="utf-8"))
    rows = prospect_rows(prospects)
    credentials = drive_credentials(client_secrets_file, token_file)
    service = build("drive", "v3", credentials=credentials)
    if file_id:
        try:
            target = service.files().get(
                fileId=file_id, fields="id,name,mimeType,webViewLink", supportsAllDrives=True
            ).execute()
        except HttpError as error:
            if error.resp.status != 404:
                raise
            target = None
    else:
        target = None
    if target is None:
        target = service.files().create(
            body={"name": SPREADSHEET_NAME, "mimeType": SHEETS_MIME},
            fields="id,name,mimeType,webViewLink",
        ).execute()
        file_id = target["id"]
        if on_create:
            on_create(file_id)
    if target["mimeType"] == SHEETS_MIME:
        sheet_id, added = update_google_sheet(credentials, file_id, rows)
        target["webViewLink"] = f"https://docs.google.com/spreadsheets/d/{file_id}/edit#gid={sheet_id}"
        return added, target
    if target["mimeType"] != XLSX_MIME:
        raise ValueError(f"Drive file {file_id} must be a Google Sheet or an existing .xlsx file")
    existing_content = service.files().get_media(
        fileId=file_id, supportsAllDrives=True
    ).execute()
    workbook, added = workbook_bytes(rows, existing_content)
    if added or SHEET_NAME.casefold() not in (
        sheet.title.casefold() for sheet in load_workbook(BytesIO(existing_content)).worksheets
    ):
        target = service.files().update(
            fileId=file_id,
            media_body=MediaIoBaseUpload(workbook, mimetype=XLSX_MIME, resumable=False),
            fields="id,name,webViewLink",
            supportsAllDrives=True,
        ).execute()
    return added, target


def _save_file_id(env_file, file_id):
    """Remember a newly created spreadsheet for later exports."""
    lines = env_file.read_text(encoding="utf-8").splitlines() if env_file.exists() else []
    setting = f"GOOGLE_DRIVE_FILE_ID={file_id}"
    for index, line in enumerate(lines):
        if line.startswith("GOOGLE_DRIVE_FILE_ID="):
            lines[index] = setting
            break
    else:
        lines.append(setting)
    env_file.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    load_dotenv(PROJECT_DIR / ".env")
    parser = argparse.ArgumentParser(description="Add new prospective jobs to a Google Sheet or Drive Excel file.")
    parser.add_argument("--input", type=Path, default=Path("prospective_jobs.json"))
    args = parser.parse_args()
    file_id = os.getenv("GOOGLE_DRIVE_FILE_ID", "").strip()
    client_secrets = os.getenv("GOOGLE_OAUTH_CLIENT_SECRETS_FILE", "").strip()
    if not client_secrets:
        parser.error("Set GOOGLE_OAUTH_CLIENT_SECRETS_FILE in .env")
    client_secrets_file = Path(client_secrets).expanduser()
    if not client_secrets_file.is_file():
        parser.error(f"OAuth client JSON not found: {client_secrets_file}")
    token_file = Path(os.getenv("GOOGLE_OAUTH_TOKEN_FILE", ".google_drive_token.json")).expanduser()
    added, updated = upload_json(
        args.input, file_id, client_secrets_file, token_file,
        on_create=lambda new_id: _save_file_id(PROJECT_DIR / ".env", new_id),
    )
    print(f"Added {added} new jobs to {updated['name']} ({updated.get('webViewLink', updated['id'])}).")
    exported = json.loads(args.input.read_text(encoding="utf-8"))
    print(f"Input jobs by source: {source_counts(row.get('source', 'Other') for row in exported)}.")


if __name__ == "__main__":
    main()
